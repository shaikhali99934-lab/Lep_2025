import openpyxl
import json
import os

def extract_karachi_data():
    excel_path = 'Zero Leprosy Project 2025 (Karachi).xlsx'
    wb = openpyxl.load_workbook(excel_path, data_only=True)
    ws = wb['Compiled']

    outcome_meta = {
        '1': {
            'title': 'Outcome 1: Strengthening Leprosy Case Detection with Interruption of Transmission through SDR-PEP',
            'shortTitle': 'Outcome 1: Strengthening Leprosy Case Detection with Interruption of Transmission through SDR-PEP',
            'color': '#3b82f6', # blue
            'icon': '🛡️',
            'desc': 'Strengthening Leprosy Case Detection with Interruption of Transmission through SDR-PEP'
        },
        '2': {
            'title': 'Outcome 2: Sustained Comprehensive Care Services',
            'shortTitle': 'Outcome 2: Sustained Comprehensive Care Services',
            'color': '#10b981', # emerald
            'icon': '🩺',
            'desc': 'Sustained Comprehensive Care Services'
        },
        '3': {
            'title': 'Outcome 3: Preparation of Leprosy Control Programme Pakistan for post-elimination phase (phase 2) has improved',
            'shortTitle': 'Outcome 3: Preparation of Leprosy Control Programme Pakistan for post-elimination phase (phase 2) has improved',
            'color': '#8b5cf6', # purple
            'icon': '🔬',
            'desc': 'Preparation of Leprosy Control Programme Pakistan for post-elimination phase (phase 2) has improved'
        },
        '4': {
            'title': 'Outcome 4: Continue support to TB Control (Prompt diagnosis and adequate treatment of TB)',
            'shortTitle': 'Outcome 4: Continue support to TB Control (Prompt diagnosis and adequate treatment of TB)',
            'color': '#f59e0b', # amber
            'icon': '🫁',
            'desc': 'Continue support to TB Control (Prompt diagnosis and adequate treatment of TB)'
        },
        '5': {
            'title': 'Outcome 5: Project Accompanying activities (Regular monitoring and planned activities from health-facility to national levels)',
            'shortTitle': 'Outcome 5: Project Accompanying activities (Regular monitoring and planned activities from health-facility to national levels)',
            'color': '#06b6d4', # cyan
            'icon': '📋',
            'desc': 'Project Accompanying activities (Regular monitoring and planned activities from health-facility to national levels)'
        }
    }

    sub_meta = {
        '1.1': 'Early Case Detection & SDR-PEP Contact Screening',
        '1.2': 'Capacity Building & Specialized Leprosy Trainings',
        '2.1': 'Clinical Consultation & Specialized Treatment',
        '2.2': 'Complication Management & Inpatient Hospitalization',
        '2.3': 'Assistive Devices, Protective Footwear & Rehabilitation',
        '2.4': 'Mental Health, Counselling & Socio-Economic Support',
        '2.5': 'IEC Awareness Campaigns & Anti-Stigma Advocacy',
        '3.1': 'Knowledge Centre & Partner Dermatologist Network',
        '3.2': 'GIS Patient Mapping & DHIS2 / Kobo Digitalization',
        '3.3': 'Biopsy Sample Collection & qPCR Laboratory Testing',
        '4.1': 'TB Contact Screening, Chest Camps & Diagnostic Clinics',
        '5.1': 'Program M&E Reviews, Webmo Tool & Annual Audits'
    }

    activities = []
    month_names = ['Jan', 'Feb', 'Mar', 'Apr', 'May', 'Jun', 'Jul', 'Aug', 'Sep', 'Oct', 'Nov', 'Dec']
    # Mapping of column indices in Excel (1-based):
    # Jan=4, Feb=5, Mar=6, Q1=7, Apr=8, May=9, Jun=10, Q2=11, Jul=12, Aug=13, Sep=14, Q3=15, Oct=16, Nov=17, Dec=18, Q4=19, Total=20
    col_map = {
        'Jan': 4, 'Feb': 5, 'Mar': 6, 'Q1': 7,
        'Apr': 8, 'May': 9, 'Jun': 10, 'Q2': 11,
        'Jul': 12, 'Aug': 13, 'Sep': 14, 'Q3': 15,
        'Oct': 16, 'Nov': 17, 'Dec': 18, 'Q4': 19,
        'Total': 20
    }

    for r in range(5, ws.max_row + 1):
        c3 = ws.cell(r, 3).value
        if c3 == 'Planned':
            c1 = str(ws.cell(r, 1).value or '').strip()
            c2 = str(ws.cell(r, 2).value or '').strip()
            
            p_tot = float(ws.cell(r, col_map['Total']).value or 0)
            a_tot = float(ws.cell(r + 1, col_map['Total']).value or 0)
            
            # calculate % accurately
            pct_tot = round((a_tot / p_tot * 100) if p_tot > 0 else (100.0 if a_tot > 0 else 0.0), 1)

            major = c1.split('.')[0] if '.' in c1 else '1'
            sub = '.'.join(c1.split('.')[:2]) if '.' in c1 else '1.1'

            # quarters
            q_data = {}
            for q_name in ['Q1', 'Q2', 'Q3', 'Q4']:
                c_idx = col_map[q_name]
                qp = float(ws.cell(r, c_idx).value or 0)
                qa = float(ws.cell(r + 1, c_idx).value or 0)
                qpct = round((qa / qp * 100) if qp > 0 else (100.0 if qa > 0 else 0.0), 1)
                q_data[q_name] = {'planned': qp, 'achieved': qa, 'pct': qpct}

            # months
            m_data = {}
            for m_name in month_names:
                c_idx = col_map[m_name]
                mp = float(ws.cell(r, c_idx).value or 0)
                ma = float(ws.cell(r + 1, c_idx).value or 0)
                mpct = round((ma / mp * 100) if mp > 0 else (100.0 if ma > 0 else 0.0), 1)
                m_data[m_name] = {'planned': mp, 'achieved': ma, 'pct': mpct}

            # Determine Health Status per User Rules:
            # - Fully Achieved: >= 70%
            # - Partially Achieved: 50% - 69.9%
            # - Target Not Achieved: < 50%
            if pct_tot >= 70.0:
                status_code = 'fully_achieved'
                status_label = 'Fully Achieved'
                status_color = 'emerald'
            elif pct_tot >= 50.0:
                status_code = 'partially_achieved'
                status_label = 'Partially Achieved'
                status_color = 'amber'
            else:
                status_code = 'not_achieved'
                status_label = 'Target Not Achieved'
                status_color = 'rose'

            activities.append({
                'id': f'act_{len(activities)+1}',
                'code': c1,
                'title': c2,
                'outcome': major,
                'outcomeTitle': outcome_meta.get(major, {}).get('title', f'Outcome {major}'),
                'outcomeShort': outcome_meta.get(major, {}).get('shortTitle', f'Outcome {major}'),
                'outcomeIcon': outcome_meta.get(major, {}).get('icon', '📌'),
                'sub': sub,
                'subTitle': sub_meta.get(sub, f'Sub-outcome {sub}'),
                'planned_total': p_tot,
                'achieved_total': a_tot,
                'pct_total': pct_tot,
                'gap': round(p_tot - a_tot, 1),
                'status_code': status_code,
                'status_label': status_label,
                'status_color': status_color,
                'quarters': q_data,
                'months': m_data
            })

    # Global summaries
    total_planned = sum(a['planned_total'] for a in activities)
    total_achieved = sum(a['achieved_total'] for a in activities)
    overall_pct = round((total_achieved / total_planned * 100) if total_planned > 0 else 0, 1)

    status_counts = {
        'fully_achieved': sum(1 for a in activities if a['status_code'] == 'fully_achieved'),
        'partially_achieved': sum(1 for a in activities if a['status_code'] == 'partially_achieved'),
        'not_achieved': sum(1 for a in activities if a['status_code'] == 'not_achieved'),
    }

    # Quarterly totals
    quarters_summary = {}
    for q in ['Q1', 'Q2', 'Q3', 'Q4']:
        qp = sum(a['quarters'][q]['planned'] for a in activities)
        qa = sum(a['quarters'][q]['achieved'] for a in activities)
        quarters_summary[q] = {
            'planned': qp,
            'achieved': qa,
            'pct': round((qa / qp * 100) if qp > 0 else 0, 1)
        }

    # Outcome totals
    outcomes_summary = {}
    for k in sorted(outcome_meta.keys()):
        acts_in_k = [a for a in activities if a['outcome'] == k]
        op = sum(a['planned_total'] for a in acts_in_k)
        oa = sum(a['achieved_total'] for a in acts_in_k)
        outcomes_summary[k] = {
            'outcome': k,
            'meta': outcome_meta[k],
            'count': len(acts_in_k),
            'planned': op,
            'achieved': oa,
            'pct': round((oa / op * 100) if op > 0 else 0, 1),
            'status_counts': {
                'fully_achieved': sum(1 for a in acts_in_k if a['status_code'] == 'fully_achieved'),
                'partially_achieved': sum(1 for a in acts_in_k if a['status_code'] == 'partially_achieved'),
                'not_achieved': sum(1 for a in acts_in_k if a['status_code'] == 'not_achieved'),
            }
        }

    # Sub-category totals
    sub_summary = {}
    for sub_k in sorted(sub_meta.keys()):
        acts_in_sub = [a for a in activities if a['sub'] == sub_k]
        sp = sum(a['planned_total'] for a in acts_in_sub)
        sa = sum(a['achieved_total'] for a in acts_in_sub)
        sub_summary[sub_k] = {
            'sub': sub_k,
            'title': sub_meta[sub_k],
            'count': len(acts_in_sub),
            'planned': sp,
            'achieved': sa,
            'pct': round((sa / sp * 100) if sp > 0 else 0, 1)
        }

    dataset = {
        'region': 'Karachi',
        'organization': 'Marie Adelaide Leprosy Centre, Karachi',
        'project': 'ZERO Leprosy Project 2025',
        'sheet': 'Compiled',
        'summary': {
            'totalActivities': len(activities),
            'totalPlanned': total_planned,
            'totalAchieved': total_achieved,
            'overallPct': overall_pct,
            'statusCounts': status_counts,
            'quarters': quarters_summary,
            'outcomes': outcomes_summary,
            'subCategories': sub_summary
        },
        'activities': activities
    }

    # Write as JS object
    js_content = f"// Zero Leprosy Project 2025 - Karachi Analytical Data\nconst KARACHI_DATA = {json.dumps(dataset, indent=2)};\n"
    with open('data_karachi.js', 'w', encoding='utf-8') as f:
        f.write(js_content)
    
    print(f"Successfully generated data_karachi.js with {len(activities)} activities.")
    print(f"Total Planned: {total_planned:,.0f} | Total Achieved: {total_achieved:,.1f} | Overall: {overall_pct}%")

if __name__ == '__main__':
    extract_karachi_data()
